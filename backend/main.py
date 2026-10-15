"""CareerSage API: single-service MVP with SQLite, secure password hashing and local resume NLP."""
from __future__ import annotations
import hashlib, hmac, json, os, re, secrets, sqlite3, time
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional
from urllib.parse import urlencode, urlparse
from urllib.request import Request, urlopen

from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = Path(os.getenv("CAREERSAGE_DATA_DIR", ROOT / "data"))
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = DATA_DIR / "careersage.sqlite3"
MAX_BYTES = 10 * 1024 * 1024
TOKEN_TTL = 60 * 60 * 24 * 7

app = FastAPI(title="CareerSage API", version="1.0.0", description="Career preparation and resume analysis MVP")
app.add_middleware(CORSMiddleware, allow_origins=os.getenv("CORS_ORIGINS", "*").split(","), allow_methods=["*"], allow_headers=["*"])

@contextmanager
def connect():
    c = sqlite3.connect(DB_PATH, timeout=15)
    c.row_factory = sqlite3.Row
    try:
        yield c
        c.commit()
    except Exception:
        c.rollback()
        raise
    finally:
        c.close()

def init_db():
    with connect() as c:
        c.execute("PRAGMA journal_mode=WAL")
        c.executescript("""
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT NOT NULL, email TEXT UNIQUE NOT NULL, phone TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS sessions(token_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), expires_at INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS resumes(id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), filename TEXT NOT NULL, extracted_text TEXT NOT NULL, analysis_json TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS answers(id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), question TEXT NOT NULL, answer TEXT NOT NULL, score INTEGER NOT NULL, feedback TEXT NOT NULL, created_at TEXT NOT NULL, time_taken_seconds REAL NOT NULL DEFAULT 0);
        CREATE TABLE IF NOT EXISTS linkedin_profiles(user_id INTEGER PRIMARY KEY REFERENCES users(id), profile_json TEXT NOT NULL, connected_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS linkedin_oauth_states(state_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), expires_at INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS job_matches(id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), company TEXT NOT NULL, title TEXT NOT NULL, source_url TEXT NOT NULL DEFAULT '', result_json TEXT NOT NULL, created_at TEXT NOT NULL);
        """)
        columns={row["name"] for row in c.execute("PRAGMA table_info(answers)")}
        if "time_taken_seconds" not in columns:
            c.execute("ALTER TABLE answers ADD COLUMN time_taken_seconds REAL NOT NULL DEFAULT 0")
init_db()

def hash_password(password: str, salt: Optional[bytes] = None) -> str:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 310_000)
    return f"{salt.hex()}:{digest.hex()}"

def verify_password(password: str, stored: str) -> bool:
    try:
        salt, expected = stored.split(":", 1)
        actual = hash_password(password, bytes.fromhex(salt)).split(":", 1)[1]
        return hmac.compare_digest(actual, expected)
    except (ValueError, TypeError): return False

def issue_session(user_id: int) -> str:
    token = secrets.token_urlsafe(32)
    with connect() as c:
        c.execute("INSERT INTO sessions VALUES(?,?,?)", (hashlib.sha256(token.encode()).hexdigest(), user_id, int(time.time()) + TOKEN_TTL))
    return token

def current_user(authorization: Optional[str] = Header(default=None)):
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(401, "Please sign in to continue.")
    token = authorization.split(" ", 1)[1]
    with connect() as c:
        row = c.execute("SELECT u.* FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token_hash=? AND s.expires_at>?", (hashlib.sha256(token.encode()).hexdigest(), int(time.time()))).fetchone()
    if not row: raise HTTPException(401, "Session expired. Please sign in again.")
    return dict(row)

def public_user(u): return {"username":u["username"],"email":u["email"],"phone":u["phone"]}

class Register(BaseModel):
    username: str
    email: str
    phone: str
    password: str

class Login(BaseModel):
    identifier: str
    password: str

@app.get("/api/health")
def health(): return {"status":"ok","service":"careersage-api","version":"1.0.0"}

@app.post("/api/auth/register")
def register(body: Register):
    if len(body.username.strip()) < 2 or len(body.password) < 8 or not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+",body.email.strip()) or not re.fullmatch(r"[+0-9 ()-]{8,18}", body.phone.strip()):
        raise HTTPException(422, "Use a 2-character name, valid phone number and password of at least 8 characters.")
    try:
        with connect() as c:
            cur = c.execute("INSERT INTO users(username,email,phone,password_hash,created_at) VALUES(?,?,?,?,?)", (body.username.strip(),body.email.lower().strip(),body.phone.strip(),hash_password(body.password),datetime.now(timezone.utc).isoformat()))
            uid=cur.lastrowid
    except sqlite3.IntegrityError: raise HTTPException(409, "That email or phone number is already registered.")
    return {"token":issue_session(uid),"user":{"username":body.username.strip(),"email":body.email.lower().strip(),"phone":body.phone.strip()}}

@app.post("/api/auth/login")
def login(body: Login):
    identifier=body.identifier.strip()
    with connect() as c:
        row=c.execute("SELECT * FROM users WHERE lower(email)=lower(?) OR phone=? OR lower(username)=lower(?)",(identifier,identifier,identifier)).fetchone()
        if not row and re.search(r"^[+0-9 ()-]+$",identifier):
            phone_digits=re.sub(r"\D","",identifier)
            row=next((candidate for candidate in c.execute("SELECT * FROM users") if re.sub(r"\D","",candidate["phone"])==phone_digits),None)
    if not row or not verify_password(body.password,row["password_hash"]): raise HTTPException(401,"Email/phone or password is incorrect.")
    return {"token":issue_session(row["id"]),"user":public_user(row)}

@app.post("/api/auth/logout")
def logout(user=Depends(current_user), authorization: str=Header()):
    token=authorization.split(" ",1)[1]
    with connect() as c: c.execute("DELETE FROM sessions WHERE token_hash=?",(hashlib.sha256(token.encode()).hexdigest(),))
    return {"ok":True}

@app.get("/api/auth/me")
def me(user=Depends(current_user)): return {"user":public_user(user)}

SKILLS = ["python","javascript","typescript","java","c++","sql","react","node.js","html","css","rest api","git","docker","aws","azure","gcp","machine learning","data analysis","excel","communication","leadership","problem solving","agile","testing","linux","kubernetes","pandas","tensorflow","figma"]
ROLES = [
    {"role":"Software Engineer","skills":["python","java","javascript","react","sql","git","testing","rest api"]},
    {"role":"Data Analyst","skills":["python","sql","data analysis","excel","pandas","communication"]},
    {"role":"Backend Developer","skills":["python","java","node.js","sql","rest api","docker","aws","git"]},
    {"role":"Frontend Developer","skills":["javascript","typescript","react","html","css","figma","git"]},
    {"role":"ML Engineer","skills":["python","machine learning","pandas","tensorflow","sql","aws"]},
]
def extract_github_links(text: str, supplied: str = "") -> list[str]:
    candidates=re.findall(r"https?://(?:www\.)?github\.com/[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+(?:\.git)?",text,re.I)
    candidates+=re.split(r"[\s,;]+",supplied.strip()) if supplied.strip() else []
    links=[]
    for candidate in candidates:
        candidate=candidate.rstrip(".,);]}")
        parsed=urlparse(candidate)
        parts=[p for p in parsed.path.strip("/").split("/") if p]
        if parsed.scheme in {"http","https"} and parsed.netloc.lower() in {"github.com","www.github.com"} and len(parts)==2:
            link=f"https://github.com/{parts[0]}/{parts[1].removesuffix('.git')}"
            if link not in links: links.append(link)
        if len(links)>=10: break
    return links

def check_resume_writing(text: str):
    words=re.findall(r"\b[\w+#'.-]+\b",text)
    letters=len(re.findall(r"[A-Za-z]",text))
    visible=len(re.sub(r"\s","",text))
    alpha_ratio=letters/max(visible,1)
    sentences=[part for part in re.split(r"[.!?]+(?:\s|$)",text) if re.search(r"[A-Za-z]",part)]
    average_sentence=round(sum(len(re.findall(r"\b[\w+#'.-]+\b",part)) for part in sentences)/max(len(sentences),1),1)
    if len(words)<20 or alpha_ratio<.45:
        read_status="Review needed"
        read_message="We extracted text, but there may not be enough clear resume text to review reliably. Check that the file is not scanned or image-only."
    else:
        read_status="Readable"
        read_message="Text was extracted successfully and is readable for the resume analysis."

    consistency=[]
    date_formats=[]
    if re.search(r"\b(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\.?\s+\d{4}\b",text,re.I): date_formats.append("month-year")
    if re.search(r"\b(?:0?[1-9]|1[0-2])[/-]\d{4}\b",text): date_formats.append("numeric-month-year")
    if len(date_formats)>1: consistency.append("Date formats appear mixed; use one style throughout.")
    bullets=[line.strip() for line in text.splitlines() if re.match(r"\s*(?:[-*•▪]|\d+[.)])\s+",line)]
    endings={bool(re.search(r"[.!?]$",line.strip())) for line in bullets if re.search(r"[A-Za-z]",line)}
    if len(endings)>1: consistency.append("Some bullet points end with punctuation and others do not.")
    repeated=re.search(r"\b([A-Za-z]{3,})\s+\1\b",text,re.I)
    if repeated: consistency.append(f"Repeated word found: “{repeated.group(1)}”.")

    grammar=[]
    grammar_patterns=[
      (r"\bI\s+has\b", "Check subject–verb agreement: “I has”."),
      (r"\b(?:he|she|it)\s+have\b", "Check subject–verb agreement: singular subjects usually take “has”."),
      (r"\b(?:they|we)\s+is\b", "Check subject–verb agreement: “they/we is”."),
      (r"\b(?:a|an)\s+(?:information|advice|feedbacks)\b", "Check article and noun agreement in this phrase."),
      (r"\s+[,.;:!?]", "Check for spaces before punctuation."),
      (r"[!?.,]{2,}", "Check repeated punctuation."),
    ]
    for pattern,message in grammar_patterns:
        if re.search(pattern,text,re.I): grammar.append(message)
    consistency=consistency[:4]; grammar=grammar[:4]
    return {
      "readability":{"status":read_status,"message":read_message,"wordCount":len(words),"averageSentenceWords":average_sentence},
      "consistency":{"status":"Review suggested" if consistency else "Looks consistent","issues":consistency},
      "grammar":{"status":"Review suggested" if grammar else "No common issues found","issues":grammar,"note":"These are basic automated checks, not a complete grammar proofread."}
    }

def analyze_text(text: str):
    normalized=re.sub(r"\s+"," ",text.lower())
    found=[s for s in SKILLS if re.search(r"(?<![a-z0-9+#])"+re.escape(s)+r"(?![a-z0-9+#])",normalized)]
    # A tiny, local TF-IDF cosine matcher keeps role ranking explainable and key-free.
    # Each role's skill profile is a reference document; the uploaded resume is the query.
    tokens=lambda value: re.findall(r"[a-z][a-z0-9+#.]*",value.lower())
    docs=[tokens(" ".join(r["skills"])) for r in ROLES]
    query=tokens(" ".join(found))
    all_docs=docs+[query]
    vocab=set(t for d in all_docs for t in d)
    import math
    idf={t:math.log((1+len(all_docs))/(1+sum(t in d for d in all_docs)))+1 for t in vocab}
    def vector(doc):
        counts={t:doc.count(t) for t in set(doc)}
        return {t:(1+math.log(n))*idf[t] for t,n in counts.items()}
    qv=vector(query); qnorm=math.sqrt(sum(v*v for v in qv.values())) or 1
    role_matches=[]
    for role,doc in zip(ROLES,docs):
        overlap=[s for s in role["skills"] if s in found]
        rv=vector(doc); rnorm=math.sqrt(sum(v*v for v in rv.values())) or 1
        cosine=sum(v*qv.get(t,0) for t,v in rv.items())/(rnorm*qnorm)
        coverage=len(overlap)/len(role["skills"])
        score=round(100*(.75*coverage+.25*cosine))
        role_matches.append({"role":role["role"],"score":score,"skills":", ".join(overlap) or "Build relevant skills"})
    role_matches.sort(key=lambda x:x["score"],reverse=True)
    return {"score":min(95,max(25,round(35+len(found)*5))),"atsScore":min(95,max(25,round(40+len(found)*5))),"contentQuality":72 if len(text)>500 else 55,"skillRelevance":min(95,45+len(found)*5),"roleAlignment":role_matches[0]["score"] if role_matches else 0,"clarity":70 if len(text.split())>100 else 55,"writingChecks":check_resume_writing(text),
      "insights":[f"Identified {len(found)} skills from the resume text.","Role scores are based on skill overlap and should guide, not replace, human review.","Add measurable outcomes and tailor keywords to the target role."],"technicalSkills":found,"softSkills":[s for s in found if s in ["communication","leadership","problem solving"]],"otherSkills":[s for s in found if s in ["agile","figma","excel"]],"matchedKeywords":found,"missingKeywords":[s for s in ["docker","aws","testing","communication"] if s not in found],"suggestedKeywords":[s for s in ["docker","aws","testing","rest api"] if s not in found][:3],"roleAlignmentList":role_matches[:3],"suggestions":["Add two quantified achievements to your project section.","Move your strongest technical skills closer to the top.","Tailor the summary and skills to your target role."]}

def extract_text(filename: str, raw: bytes):
    ext=Path(filename).suffix.lower()
    if ext==".pdf":
        from pypdf import PdfReader
        import io
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages)
    if ext==".docx":
        from docx import Document
        import io
        doc=Document(io.BytesIO(raw))
        return "\n".join(p.text for p in doc.paragraphs)+"\n"+"\n".join(" ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
    raise HTTPException(415,"Upload a PDF or DOCX resume. Legacy DOC files are not supported for server analysis.")

@app.post("/api/resumes")
async def upload_resume(file: UploadFile=File(...), github_links: str=Form(default=""), user=Depends(current_user)):
    name=Path(file.filename or "resume").name
    if Path(name).suffix.lower() not in {".pdf",".docx"}: raise HTTPException(415,"Upload a PDF or DOCX resume.")
    raw=await file.read(MAX_BYTES+1)
    if not raw or len(raw)>MAX_BYTES: raise HTTPException(413,"Resume must be non-empty and no larger than 10 MB.")
    try: text=extract_text(name,raw)
    except HTTPException: raise
    except Exception as e: raise HTTPException(422,"We could not read this file. Please upload a valid PDF or DOCX.") from e
    if len(text.strip())<40: raise HTTPException(422,"This resume has too little selectable text to analyze. Try an accessible-text PDF or DOCX.")
    analysis=analyze_text(text)
    analysis["githubLinks"]=extract_github_links(text,github_links)
    with connect() as c:
        cur=c.execute("INSERT INTO resumes(user_id,filename,extracted_text,analysis_json,created_at) VALUES(?,?,?,?,?)",(user["id"],name,text,json.dumps(analysis),datetime.now(timezone.utc).isoformat()))
        rid=cur.lastrowid
    return {"id":rid,"fileName":name,"status":"complete","analysis":analysis}

@app.get("/api/resumes/latest/analysis")
def latest_analysis(user=Depends(current_user)):
    with connect() as c: row=c.execute("SELECT id,filename,extracted_text,analysis_json,created_at FROM resumes WHERE user_id=? ORDER BY id DESC LIMIT 1",(user["id"],)).fetchone()
    if not row: raise HTTPException(404,"No resume has been analyzed yet.")
    analysis=json.loads(row["analysis_json"])
    analysis.setdefault("writingChecks",check_resume_writing(row["extracted_text"]))
    return {"id":row["id"],"fileName":row["filename"],"createdAt":row["created_at"],"analysis":analysis}

def linkedin_config():
    return (os.getenv("LINKEDIN_CLIENT_ID", "").strip(), os.getenv("LINKEDIN_CLIENT_SECRET", "").strip(), os.getenv("LINKEDIN_REDIRECT_URI", "http://127.0.0.1:8000/api/linkedin/callback").strip())

@app.get("/api/linkedin/status")
def linkedin_status(user=Depends(current_user)):
    client_id,client_secret,redirect_uri=linkedin_config()
    with connect() as c:
        row=c.execute("SELECT profile_json,connected_at FROM linkedin_profiles WHERE user_id=?",(user["id"],)).fetchone()
    return {"configured":bool(client_id and client_secret),"connected":bool(row),"profile":json.loads(row["profile_json"]) if row else None,"connectedAt":row["connected_at"] if row else None}

@app.post("/api/linkedin/connect")
def linkedin_connect(user=Depends(current_user)):
    client_id,client_secret,redirect_uri=linkedin_config()
    if not client_id or not client_secret:
        raise HTTPException(503,"LinkedIn connection is not configured yet. Add LINKEDIN_CLIENT_ID and LINKEDIN_CLIENT_SECRET to the backend environment.")
    state=secrets.token_urlsafe(32)
    with connect() as c:
        c.execute("DELETE FROM linkedin_oauth_states WHERE expires_at<=?",(int(time.time()),))
        c.execute("INSERT INTO linkedin_oauth_states VALUES(?,?,?)",(hashlib.sha256(state.encode()).hexdigest(),user["id"],int(time.time())+600))
    query=urlencode({"response_type":"code","client_id":client_id,"redirect_uri":redirect_uri,"state":state,"scope":"openid profile email"})
    return {"authorizationUrl":"https://www.linkedin.com/oauth/v2/authorization?"+query}

@app.get("/api/linkedin/callback")
def linkedin_callback(code: str="", state: str="", error: str=""):
    frontend=os.getenv("CAREERSAGE_FRONTEND_URL", "/")
    if error or not code or not state:
        return RedirectResponse(frontend+"?linkedin=cancelled",status_code=303)
    state_hash=hashlib.sha256(state.encode()).hexdigest()
    with connect() as c:
        row=c.execute("SELECT user_id,expires_at FROM linkedin_oauth_states WHERE state_hash=?",(state_hash,)).fetchone()
        c.execute("DELETE FROM linkedin_oauth_states WHERE state_hash=?",(state_hash,))
    if not row or row["expires_at"]<int(time.time()):
        return RedirectResponse(frontend+"?linkedin=failed",status_code=303)
    client_id,client_secret,redirect_uri=linkedin_config()
    try:
        body=urlencode({"grant_type":"authorization_code","code":code,"client_id":client_id,"client_secret":client_secret,"redirect_uri":redirect_uri}).encode()
        request=Request("https://www.linkedin.com/oauth/v2/accessToken",data=body,headers={"Content-Type":"application/x-www-form-urlencoded"})
        with urlopen(request,timeout=12) as response: token_data=json.loads(response.read())
        request=Request("https://api.linkedin.com/v2/userinfo",headers={"Authorization":"Bearer "+token_data["access_token"],"Accept":"application/json"})
        with urlopen(request,timeout=12) as response: raw_profile=json.loads(response.read())
        profile={key:raw_profile[key] for key in ("sub","name","given_name","family_name","picture","email") if raw_profile.get(key)}
        if not profile.get("sub"): raise ValueError("LinkedIn profile identifier missing")
        connected_at=datetime.now(timezone.utc).isoformat()
        with connect() as c:
            c.execute("INSERT INTO linkedin_profiles(user_id,profile_json,connected_at) VALUES(?,?,?) ON CONFLICT(user_id) DO UPDATE SET profile_json=excluded.profile_json,connected_at=excluded.connected_at",(row["user_id"],json.dumps(profile),connected_at))
        return RedirectResponse(frontend+"?linkedin=connected",status_code=303)
    except Exception:
        return RedirectResponse(frontend+"?linkedin=failed",status_code=303)

@app.delete("/api/linkedin/disconnect")
def linkedin_disconnect(user=Depends(current_user)):
    with connect() as c: c.execute("DELETE FROM linkedin_profiles WHERE user_id=?",(user["id"],))
    return {"ok":True}

class JobMatchRequest(BaseModel):
    company: str = Field(min_length=1,max_length=120)
    title: str = Field(min_length=1,max_length=120)
    description: str = Field(min_length=30,max_length=20000)
    source_url: str = Field(default="",max_length=500)

@app.post("/api/matches/jobs")
def match_job(body: JobMatchRequest,user=Depends(current_user)):
    if body.source_url:
        parsed=urlparse(body.source_url)
        if parsed.scheme!="https" or parsed.netloc.lower() not in {"linkedin.com","www.linkedin.com"} or not parsed.path.startswith("/jobs/"):
            raise HTTPException(422,"Use a LinkedIn job URL starting with https://www.linkedin.com/jobs/.")
    with connect() as c:
        resume=c.execute("SELECT extracted_text FROM resumes WHERE user_id=? ORDER BY id DESC LIMIT 1",(user["id"],)).fetchone()
    if not resume: raise HTTPException(404,"Upload a resume before matching jobs.")
    resume_skills=[s for s in SKILLS if re.search(r"(?<![a-z0-9+#])"+re.escape(s)+r"(?![a-z0-9+#])",resume["extracted_text"].lower())]
    job_skills=[s for s in SKILLS if re.search(r"(?<![a-z0-9+#])"+re.escape(s)+r"(?![a-z0-9+#])",body.description.lower())]
    matched=[s for s in job_skills if s in resume_skills]
    missing=[s for s in job_skills if s not in resume_skills]
    score=round(100*len(matched)/len(job_skills)) if job_skills else 0
    result={"company":body.company.strip(),"title":body.title.strip(),"score":score,"matchedSkills":matched,"missingSkills":missing,"jobSkills":job_skills,"sourceUrl":body.source_url,"note":"This match is based on skill keyword overlap in your latest resume and the job description you supplied. LinkedIn job listings are not fetched automatically."}
    created=datetime.now(timezone.utc).isoformat()
    with connect() as c:
        c.execute("INSERT INTO job_matches(user_id,company,title,source_url,result_json,created_at) VALUES(?,?,?,?,?,?)",(user["id"],result["company"],result["title"],body.source_url,json.dumps(result),created))
    return result

@app.get("/api/matches/jobs")
def list_job_matches(user=Depends(current_user)):
    with connect() as c:
        rows=c.execute("SELECT result_json,created_at FROM job_matches WHERE user_id=? ORDER BY id DESC LIMIT 10",(user["id"],)).fetchall()
    return [{**json.loads(row["result_json"]),"createdAt":row["created_at"]} for row in rows]

@app.get("/api/dashboard")
def dashboard(user=Depends(current_user)):
    with connect() as c:
        row=c.execute("SELECT filename,analysis_json FROM resumes WHERE user_id=? ORDER BY id DESC LIMIT 1",(user["id"],)).fetchone()
        count=c.execute("SELECT count(*) FROM answers WHERE user_id=?",(user["id"],)).fetchone()[0]
    analysis=json.loads(row["analysis_json"]) if row else analyze_text("")
    return {"user":public_user(user),"resume":{"fileName":row["filename"],**analysis} if row else None,"dashboard":{"readiness":analysis["score"],"skillsIdentified":len(analysis["technicalSkills"]),"recommendedRoles":len(analysis["roleAlignmentList"]),"preparationStatus":min(100,count*10),"skillAnalysis":[[s] for s in analysis["technicalSkills"]],"roleMatches":[[r["role"],r["score"]] for r in analysis["roleAlignmentList"]],"strengths":analysis["technicalSkills"][:3],"areasToImprove":analysis["missingKeywords"][:3],"nextSteps":analysis["suggestions"]}}

@app.get("/api/interview/questions")
def questions(category: str="", company: str="", difficulty: str="", topic: str=""):
    with open(ROOT/"backend"/"questions.json",encoding="utf-8") as f: items=json.load(f)
    return [q for q in items if (not category or q["category"].lower()==category.lower()) and (not company or q["company"].lower()==company.lower()) and (not difficulty or q["difficulty"].lower()==difficulty.lower()) and (not topic or q["topic"].lower()==topic.lower())]

class Answer(BaseModel):
    question: str
    answer: str
    time_taken_seconds: float = Field(default=0, ge=0, le=86400)
@app.post("/api/interview/answers")
def submit_answer(body: Answer,user=Depends(current_user)):
    words=len(body.answer.split()); score=min(100,max(20,35+min(words,100)//3+(15 if any(w in body.answer.lower() for w in ["result","impact","learned","because","measured"]) else 0)))
    feedback="Add more specific details and explain your role." if words<30 else "Good structure. Strengthen it with a measurable result and what you learned."
    with connect() as c: c.execute("INSERT INTO answers(user_id,question,answer,score,feedback,created_at,time_taken_seconds) VALUES(?,?,?,?,?,?,?)",(user["id"],body.question,body.answer,score,feedback,datetime.now(timezone.utc).isoformat(),body.time_taken_seconds))
    return {"score":score,"feedback":feedback,"time_taken_seconds":body.time_taken_seconds}

@app.get("/api/progress")
def progress(user=Depends(current_user)):
    with connect() as c:
        rows=c.execute("SELECT score,question,created_at,time_taken_seconds FROM answers WHERE user_id=? ORDER BY id DESC LIMIT 12",(user["id"],)).fetchall()
        totals=c.execute("SELECT COUNT(*),AVG(score),AVG(NULLIF(time_taken_seconds,0)) FROM answers WHERE user_id=?",(user["id"],)).fetchone()
    vals=[dict(week=f"Practice {len(rows)-i}",accuracy=r["score"],attempted=i+1,time=round(r["time_taken_seconds"],1)) for i,r in enumerate(reversed(rows))]
    activity=[[r["question"][:90],"Mixed",1,f'{r["score"]}%',r["created_at"][:10],r["time_taken_seconds"]] for r in rows]
    return {"progress":vals,"recentActivity":activity,"summary":{"answers":totals[0],"averageScore":round(totals[1] or 0),"averageTimeSeconds":round(totals[2],1) if totals[2] is not None else None}}

if (ROOT/"dist").exists():
    app.mount("/assets",StaticFiles(directory=ROOT/"dist"/"assets"),name="assets")
    @app.get("/{path:path}")
    def frontend(path: str):
        candidate=(ROOT/"dist"/path).resolve()
        if candidate.is_file() and ROOT/"dist" in candidate.parents: return FileResponse(candidate)
        return FileResponse(ROOT/"dist"/"index.html")
